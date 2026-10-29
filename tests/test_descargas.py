import io
import re
from datetime import time
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.extensions import db
from app.models import Entrenamiento
from app.utils.descargas import guardar_buffer


def _comprobar_pdf(datos):
    assert datos.startswith(b"%PDF-")
    assert b"%%EOF" in datos[-1024:]


def _comprobar_excel(datos):
    libro = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    assert libro.sheetnames
    hoja = libro[libro.sheetnames[0]]
    assert hoja.max_row >= 1
    assert hoja.max_column >= 1
    libro.close()


def test_documentos_se_guardan_bajo_una_raiz_y_por_tipo(app):
    with app.app_context():
        recibo = guardar_buffer(
            io.BytesIO(b"recibo"), "recibos", "recibo_000001.pdf"
        )
        reporte = guardar_buffer(
            io.BytesIO(b"reporte"), "reportes_pagos", "pagos.xlsx"
        )
        ficha = guardar_buffer(
            io.BytesIO(b"ficha"), "perfiles", "ficha_alumno.pdf"
        )

        raiz = Path(app.config["DOWNLOAD_FOLDER"]).resolve()
        assert recibo == raiz / "Pagos" / "Recibos" / "recibo_000001.pdf"
        assert reporte == raiz / "Pagos" / "Reportes" / "pagos.xlsx"
        assert ficha == raiz / "Alumnos" / "Perfiles" / "ficha_alumno.pdf"
        assert recibo.read_bytes() == b"recibo"
        assert reporte.read_bytes() == b"reporte"
        assert ficha.read_bytes() == b"ficha"


def test_nombre_repetido_no_sobrescribe_documento_anterior(app):
    with app.app_context():
        primero = guardar_buffer(io.BytesIO(b"uno"), "alumnos", "alumnos.pdf")
        segundo = guardar_buffer(io.BytesIO(b"dos"), "alumnos", "alumnos.pdf")

        assert primero != segundo
        assert primero.read_bytes() == b"uno"
        assert segundo.read_bytes() == b"dos"


def test_exportaciones_de_calendario_y_asistencia_son_validas(client, app):
    mensual = client.get("/calendario/exportar-mensual.pdf?anio=2026&mes=8")
    semanal = client.get("/calendario/exportar-semanal.pdf?fecha=2026-08-17")
    _comprobar_pdf(mensual.data)
    _comprobar_pdf(semanal.data)

    with app.app_context():
        entrenamiento = Entrenamiento(dia_semana=0, hora_inicio=time(18, 0), lugar="Cancha principal")
        db.session.add(entrenamiento)
        db.session.commit()
        id_entrenamiento = entrenamiento.id_entrenamiento

    parametros = [("fecha", "2026-08-17"), ("columnas", "alumno"), ("columnas", "estado")]
    pdf = client.get(
        f"/calendario/entrenamientos/{id_entrenamiento}/asistencia/exportar.pdf",
        query_string=parametros,
    )
    excel = client.get(
        f"/calendario/entrenamientos/{id_entrenamiento}/asistencia/exportar.xlsx",
        query_string=parametros,
    )
    assert pdf.status_code == excel.status_code == 200
    _comprobar_pdf(pdf.data)
    _comprobar_excel(excel.data)

