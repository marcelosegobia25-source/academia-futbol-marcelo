from datetime import date, time, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError

from app.extensions import db
from app.models import (
    Alumno, AnulacionPago, Auditoria, Campeonato, Categoria, CierreCaja, Configuracion,
    Equipo, EventoAgenda, HistorialAlumno, Pago, Partido, PodioCampeonato,
    Entrenamiento, SesionEntrenamiento, AsistenciaEntrenamiento,
)
from app.utils.excel import exportar_tabla_excel
from app.utils.reglas import siguiente_mes_a_pagar


def alumno_base(ci="MEJORA-1", estado="Activo"):
    categoria = Categoria.query.filter_by(nombre="13").first()
    alumno = Alumno(
        nombre="María", apellido="Control", fecha_nacimiento=date(2013, 5, 1),
        ci=ci, estado=estado, id_categoria=categoria.id_categoria,
    )
    db.session.add(alumno)
    db.session.commit()
    return alumno, categoria


def campeonato_con_equipos(cantidad):
    categoria = Categoria.query.filter_by(nombre="11").first()
    equipos = [
        Equipo(nombre=f"Fixture {cantidad}-{indice}", id_categoria=categoria.id_categoria)
        for indice in range(cantidad)
    ]
    campeonato = Campeonato(
        nombre=f"Copa {cantidad}", fecha_inicio=date(2026, 8, 1),
        categorias=[categoria], equipos=equipos,
    )
    db.session.add(campeonato)
    db.session.commit()
    return campeonato, categoria, equipos


@pytest.mark.parametrize("cantidad", [2, 3, 4, 5, 6, 7, 8, 9])
def test_fixture_todos_contra_todos_cantidad_correcta(client, cantidad):
    campeonato, _, _ = campeonato_con_equipos(cantidad)
    respuesta = client.post(f"/campeonatos/{campeonato.id_campeonato}/generar-fixture")
    assert respuesta.status_code == 302
    assert Partido.query.filter_by(
        id_campeonato=campeonato.id_campeonato
    ).count() == cantidad * (cantidad - 1) // 2


@pytest.mark.parametrize("cantidad", [3, 5, 7, 9])
def test_fixture_impar_distribuye_jornadas(client, cantidad):
    campeonato, _, _ = campeonato_con_equipos(cantidad)
    client.post(f"/campeonatos/{campeonato.id_campeonato}/generar-fixture")
    partidos = Partido.query.filter_by(id_campeonato=campeonato.id_campeonato).all()
    assert len({p.jornada for p in partidos}) == cantidad
    assert max(sum(p.jornada == jornada for p in partidos) for jornada in range(1, cantidad + 1)) == cantidad // 2


def test_base_rechaza_goles_negativos(client):
    campeonato, categoria, equipos = campeonato_con_equipos(2)
    db.session.add(Partido(
        fecha=date.today(), id_equipo=equipos[0].id_equipo,
        id_equipo_rival=equipos[1].id_equipo, id_categoria=categoria.id_categoria,
        rival=equipos[1].nombre, id_campeonato=campeonato.id_campeonato,
        goles_local=-1,
    ))
    with pytest.raises(IntegrityError):
        db.session.commit()
    db.session.rollback()


def test_base_rechaza_local_igual_visitante(client):
    campeonato, categoria, equipos = campeonato_con_equipos(2)
    db.session.add(Partido(
        fecha=date.today(), id_equipo=equipos[0].id_equipo,
        id_equipo_rival=equipos[0].id_equipo, id_categoria=categoria.id_categoria,
        rival=equipos[0].nombre, id_campeonato=campeonato.id_campeonato,
    ))
    with pytest.raises(IntegrityError):
        db.session.commit()
    db.session.rollback()


@pytest.mark.parametrize("fase", ["Desconocida", "", "Octavos"])
def test_base_rechaza_fases_invalidas(client, fase):
    campeonato, categoria, equipos = campeonato_con_equipos(2)
    db.session.add(Partido(
        fecha=date.today(), id_equipo=equipos[0].id_equipo,
        id_equipo_rival=equipos[1].id_equipo, id_categoria=categoria.id_categoria,
        rival=equipos[1].nombre, id_campeonato=campeonato.id_campeonato, fase=fase,
    ))
    with pytest.raises(IntegrityError):
        db.session.commit()
    db.session.rollback()


@pytest.mark.parametrize("estado", ["Abierto", "", "Borrado"])
def test_base_rechaza_estados_campeonato_invalidos(client, estado):
    db.session.add(Campeonato(
        nombre="Inválido", fecha_inicio=date.today(), estado=estado
    ))
    with pytest.raises(IntegrityError):
        db.session.commit()
    db.session.rollback()


def test_equipo_duplicado_misma_categoria_rechazado(client):
    categoria = Categoria.query.filter_by(nombre="13").first()
    db.session.add(Equipo(nombre="Duplicado", id_categoria=categoria.id_categoria))
    db.session.commit()
    db.session.add(Equipo(nombre="duplicado", id_categoria=categoria.id_categoria))
    with pytest.raises(IntegrityError):
        db.session.commit()
    db.session.rollback()


def test_categoria_solapada_rechazada(client):
    db.session.add(Categoria(nombre="Solapada", edad_min=12, edad_max=14))
    with pytest.raises(IntegrityError):
        db.session.commit()
    db.session.rollback()


def test_eliminatoria_con_un_equipo_no_falla(client):
    campeonato, categoria, _ = campeonato_con_equipos(1)
    respuesta = client.post(
        f"/campeonatos/{campeonato.id_campeonato}/eliminatorias/generar",
        data={"id_categoria": categoria.id_categoria},
    )
    assert respuesta.status_code in {200, 302}
    assert Partido.query.filter_by(id_campeonato=campeonato.id_campeonato).count() == 0


def test_cambio_estructural_bloqueado_con_fixture(client):
    campeonato, categoria, equipos = campeonato_con_equipos(2)
    client.post(f"/campeonatos/{campeonato.id_campeonato}/generar-fixture")
    otra_categoria = Categoria.query.filter_by(nombre="13").first()
    respuesta = client.post(
        f"/campeonatos/{campeonato.id_campeonato}/editar",
        data={
            "nombre": campeonato.nombre, "fecha_inicio": "2026-08-01",
            "estado": "Planificado", "formato": "Liga", "vueltas": "1",
            "clasificados": "2", "puntos_victoria": "3",
            "puntos_empate": "1", "puntos_derrota": "0",
            "categorias": [otra_categoria.id_categoria],
            "equipos": [equipo.id_equipo for equipo in equipos],
        },
    )
    # La ruta puede redirigir por el bloqueo estructural o volver a mostrar
    # el formulario si la nueva combinación ya es incompatible. En ambos
    # casos la estructura persistida debe permanecer intacta.
    assert respuesta.status_code in {200, 302}
    assert categoria in campeonato.categorias


def test_qr_exige_numero_operacion(client):
    alumno, _ = alumno_base()
    respuesta = client.post(
        f"/pagos/alumno/{alumno.id_alumno}/nuevo",
        data={"mes_pagado": "2026-07", "fecha_pago": "2026-07-20", "modo_pago": "QR"},
    )
    assert respuesta.status_code == 200
    assert Pago.query.count() == 0


def test_qr_guarda_numero_y_auditoria(client):
    alumno, _ = alumno_base()
    client.post(
        f"/pagos/alumno/{alumno.id_alumno}/nuevo",
        data={
            "mes_pagado": "2026-07", "fecha_pago": "2026-07-20",
            "modo_pago": "QR", "numero_operacion": "QR-9981",
        },
    )
    pago = Pago.query.one()
    assert pago.numero_operacion == "QR-9981"
    assert Auditoria.query.filter_by(modulo="Pagos", accion="Registrar").count() == 1


def test_anulacion_exige_motivo(client):
    alumno, _ = alumno_base()
    pago = Pago(
        alumno=alumno, fecha_pago=date.today(), fecha_vencimiento=date.today(),
        mes_pagado="2026-07", modo_pago="Efectivo", monto=100,
    )
    db.session.add(pago)
    db.session.commit()
    client.post(f"/pagos/{pago.id_pago}/eliminar", data={"motivo": ""})
    assert pago.eliminado_en is None
    assert AnulacionPago.query.count() == 0


def test_anulacion_y_pago_correctivo_quedan_vinculados(client):
    alumno, _ = alumno_base()
    pago = Pago(
        alumno=alumno, fecha_pago=date(2026, 7, 20),
        fecha_vencimiento=date(2026, 8, 20), mes_pagado="2026-07",
        modo_pago="Efectivo", monto=100,
    )
    alumno.fecha_base_pago = date(2026, 7, 20)
    db.session.add(pago)
    db.session.commit()
    client.post(
        f"/pagos/{pago.id_pago}/eliminar",
        data={"motivo": "Modo de pago incorrecto"},
    )
    client.post(
        f"/pagos/alumno/{alumno.id_alumno}/nuevo",
        data={"mes_pagado": "2026-07", "modo_pago": "Efectivo"},
    )
    anulacion = AnulacionPago.query.one()
    assert anulacion.motivo == "Modo de pago incorrecto"
    assert anulacion.id_pago_reemplazo is not None


def test_cierre_caja_concilia_metodos(client):
    alumno, _ = alumno_base()
    db.session.add_all([
        Pago(
            alumno=alumno, fecha_pago=date.today(), fecha_vencimiento=date.today(),
            mes_pagado="2026-06", modo_pago="Efectivo", monto=80,
        ),
        Pago(
            alumno=alumno, fecha_pago=date.today(), fecha_vencimiento=date.today(),
            mes_pagado="2026-07", modo_pago="QR", numero_operacion="Q1", monto=120,
        ),
    ])
    db.session.commit()
    respuesta = client.post(
        "/pagos/caja",
        data={
            "fecha": date.today().isoformat(),
            "efectivo_contado": "80", "qr_confirmado": "120",
        },
    )
    assert respuesta.status_code == 302
    cierre = CierreCaja.query.one()
    assert cierre.diferencia == 0


def test_cambio_fecha_base_deja_historial_y_respaldo_logico(client):
    alumno, _ = alumno_base()
    alumno.fecha_base_pago = date(2026, 1, 10)
    db.session.commit()
    client.post(
        f"/pagos/alumno/{alumno.id_alumno}/fecha-base",
        data={
            "fecha_base_pago": "2026-02-10",
            "motivo": "Corrección de fecha contractual",
        },
    )
    historial = HistorialAlumno.query.filter_by(
        id_alumno=alumno.id_alumno, tipo="fecha_base_pago"
    ).one()
    assert historial.valor_anterior == "2026-01-10"
    assert historial.valor_nuevo == "2026-02-10"


@pytest.mark.parametrize("valor", ["=2+2", "+SUM(A1:A2)", "-1+2", "@cmd"])
def test_excel_neutraliza_formulas(valor):
    buffer = exportar_tabla_excel("Seguro", ["Dato"], [[valor]])
    libro = load_workbook(buffer, data_only=False)
    assert libro.active["A2"].value.startswith("'")
    assert libro.active["A2"].data_type != "f"


def test_excel_conserva_enlace_google_maps_como_hipervinculo():
    enlace = "https://maps.app.goo.gl/CanchaPrueba123"
    buffer = exportar_tabla_excel(
        "Ubicaciones", ["Lugar", "Google Maps"], [["Cancha de prueba", enlace]]
    )
    libro = load_workbook(buffer)
    assert libro.active["B2"].value == enlace
    assert libro.active["B2"].hyperlink.target == enlace


def test_reportes_campeonato_pdf_y_excel(client):
    campeonato, _, _ = campeonato_con_equipos(2)
    client.post(f"/campeonatos/{campeonato.id_campeonato}/generar-fixture")
    excel = client.get(f"/campeonatos/{campeonato.id_campeonato}/resumen.xlsx")
    pdf = client.get(f"/campeonatos/{campeonato.id_campeonato}/resumen.pdf")
    assert excel.status_code == 200 and excel.data.startswith(b"PK")
    assert pdf.status_code == 200 and pdf.data.startswith(b"%PDF")


def test_fixture_y_posiciones_se_exportan_en_pdf_y_excel(client):
    campeonato, categoria, equipos = campeonato_con_equipos(3)
    client.post(f"/campeonatos/{campeonato.id_campeonato}/generar-fixture")

    respuestas = {
        "fixture_xlsx": client.get(
            f"/campeonatos/{campeonato.id_campeonato}/fixture.xlsx",
            query_string={"id_categoria": categoria.id_categoria, "fase": "Liga"},
        ),
        "fixture_pdf": client.get(
            f"/campeonatos/{campeonato.id_campeonato}/fixture.pdf",
            query_string={"id_categoria": categoria.id_categoria, "fase": "Liga"},
        ),
        "posiciones_xlsx": client.get(
            f"/campeonatos/{campeonato.id_campeonato}/posiciones.xlsx",
            query_string={"id_categoria": categoria.id_categoria},
        ),
        "posiciones_pdf": client.get(
            f"/campeonatos/{campeonato.id_campeonato}/posiciones.pdf",
            query_string={"id_categoria": categoria.id_categoria},
        ),
    }
    assert respuestas["fixture_xlsx"].data.startswith(b"PK")
    assert respuestas["fixture_pdf"].data.startswith(b"%PDF")
    assert respuestas["posiciones_xlsx"].data.startswith(b"PK")
    assert respuestas["posiciones_pdf"].data.startswith(b"%PDF")

    libro_fixture = load_workbook(BytesIO(respuestas["fixture_xlsx"].data))
    valores_fixture = list(libro_fixture.active.values)
    assert valores_fixture[0][0] == "N°"
    assert len(valores_fixture) == 4  # encabezado + tres partidos

    libro_posiciones = load_workbook(BytesIO(respuestas["posiciones_xlsx"].data))
    valores_posiciones = list(libro_posiciones.active.values)
    assert valores_posiciones[0][3] == "Equipo"
    assert {fila[3] for fila in valores_posiciones[1:]} == {equipo.nombre for equipo in equipos}


def test_exportacion_acepta_nombre_con_caracteres_especiales(client):
    alumno_base(ci="EXP-ESPECIAL")
    parametros = {
        "nombre_archivo": "Academia / gestión: [julio]? <control & final>",
        "columnas": ["alumno", "ci", "categoria"],
    }
    excel = client.get("/alumnos/exportar", query_string=parametros)
    pdf = client.get("/alumnos/exportar.pdf", query_string=parametros)

    assert excel.status_code == 200 and excel.data.startswith(b"PK")
    assert pdf.status_code == 200 and pdf.data.startswith(b"%PDF")
    libro = load_workbook(BytesIO(excel.data))
    assert not any(caracter in libro.active.title for caracter in r"\/:*?[]")
    assert libro.active.freeze_panes == "A2"
    assert libro.active.auto_filter.ref


@pytest.mark.parametrize("ruta", [
    "/alumnos/exportar",
    "/alumnos/exportar.pdf",
    "/pagos/exportar",
    "/pagos/exportar.pdf",
])
def test_exportacion_configurable_rechaza_solicitud_sin_columnas(client, ruta):
    assert client.get(ruta).status_code == 400


def test_perfil_muestra_historial_completo_de_pagos(client):
    alumno, _ = alumno_base(ci="HISTORIAL-COMPLETO")
    for numero in range(1, 8):
        db.session.add(Pago(
            alumno=alumno,
            fecha_pago=date(2026, numero, 1),
            fecha_vencimiento=date(2026, numero, 28),
            mes_pagado=f"2026-{numero:02d}",
            modo_pago="Efectivo",
            monto=100,
        ))
    db.session.commit()

    pagina = client.get(f"/alumnos/{alumno.id_alumno}")
    assert b"Historial completo de pagos" in pagina.data
    assert b"2026-01" in pagina.data
    assert pagina.data.find(b"2026-07") < pagina.data.find(b"2026-01")


def test_exportacion_ics_fue_retirada(client):
    assert client.get("/calendario/exportar.ics").status_code == 404


def test_ficha_pdf_usa_gestion_configurada(client, app, monkeypatch):
    alumno, _ = alumno_base(ci="FICHA-GESTION")
    with app.app_context():
        config = Configuracion.obtener()
        config.gestion_ficha = 2031
        db.session.commit()
    recibido = {}

    def ficha_simulada(alumno, logo_path, foto_path=None, gestion=None, **kwargs):
        recibido["gestion"] = gestion
        return BytesIO(b"%PDF-ficha-prueba")

    monkeypatch.setattr("app.routes.alumnos.generar_ficha_alumno", ficha_simulada)

    respuesta = client.get(f"/alumnos/{alumno.id_alumno}/ficha.pdf")

    assert respuesta.status_code == 200
    assert respuesta.data.startswith(b"%PDF")
    assert recibido["gestion"] == 2031


def test_paginacion_alumnos(client):
    categoria = Categoria.query.filter_by(nombre="13").first()
    db.session.add_all([
        Alumno(
            nombre=f"Alumno {i}", apellido="Paginado",
            fecha_nacimiento=date(2013, 1, 1), ci=f"PAG-{i}",
            id_categoria=categoria.id_categoria,
        )
        for i in range(45)
    ])
    db.session.commit()
    pagina = client.get("/alumnos/?page=2")
    assert pagina.status_code == 200
    assert b"ALU-" in pagina.data


def test_migracion_actual_e_indices(app):
    versiones = db.session.execute(text("SELECT version FROM schema_migration")).scalars().all()
    assert 8 in versiones
    indices = {
        fila[1] for fila in db.session.execute(text("PRAGMA index_list('pago')")).all()
    }
    assert "ix_pago_mes_activo" in indices


def _formulario_reactivacion(alumno, categoria, fecha=None):
    datos = {
        "nombre": alumno.nombre,
        "apellido": alumno.apellido,
        "fecha_nacimiento": alumno.fecha_nacimiento.isoformat(),
        "ci": alumno.ci or "",
        "estado": "Activo",
        "id_categoria": str(categoria.id_categoria),
        "beca": "Ninguna",
        "motivo_cambio": "Retorno a entrenamientos",
    }
    if fecha:
        datos["fecha_reactivacion"] = fecha
    return datos


def test_reactivacion_exige_nueva_fecha(client):
    alumno, categoria = alumno_base(estado="Retirado")
    respuesta = client.post(
        f"/alumnos/{alumno.id_alumno}/editar",
        data=_formulario_reactivacion(alumno, categoria),
    )
    assert respuesta.status_code == 200
    assert alumno.estado == "Retirado"
    assert alumno.fecha_ultima_reactivacion is None


def test_reactivacion_abre_ciclo_y_preserva_historial(client):
    alumno, categoria = alumno_base(estado="Retirado")
    alumno.fecha_base_pago = date(2026, 1, 10)
    pago_anterior = Pago(
        alumno=alumno, fecha_pago=date(2026, 3, 10),
        fecha_vencimiento=date(2026, 4, 10), mes_pagado="2026-03",
        modo_pago="Efectivo", monto=100,
    )
    db.session.add(pago_anterior)
    db.session.commit()

    respuesta = client.post(
        f"/alumnos/{alumno.id_alumno}/editar",
        data=_formulario_reactivacion(alumno, categoria, "2026-07-15"),
    )
    assert respuesta.status_code == 302
    assert alumno.estado == "Activo"
    assert alumno.fecha_base_pago == date(2026, 7, 15)
    assert alumno.fecha_ultima_reactivacion == date(2026, 7, 15)
    assert pago_anterior.fecha_vencimiento == date(2026, 4, 10)
    assert alumno.estado_mensualidad == "Sin pagos"
    assert siguiente_mes_a_pagar(alumno) == "2026-07"
    historial = HistorialAlumno.query.filter_by(
        id_alumno=alumno.id_alumno, tipo="reactivacion"
    ).one()
    assert historial.valor_anterior == "2026-01-10"
    assert historial.valor_nuevo == "2026-07-15"


def test_retirado_tiene_estado_propio_y_no_aparece_como_deudor(client):
    retirado, categoria = alumno_base(ci="RETIRADO-DEUDA", estado="Retirado")
    retirado.nombre = "RetiradoUnico"
    retirado.fecha_retiro = date.today()
    db.session.add(Pago(
        alumno=retirado,
        fecha_pago=date.today() - timedelta(days=90),
        fecha_vencimiento=date.today() - timedelta(days=60),
        mes_pagado=(date.today() - timedelta(days=90)).strftime("%Y-%m"),
        modo_pago="Efectivo",
        monto=100,
    ))
    activo = Alumno(
        nombre="ActivoDeudorUnico", apellido="Control",
        fecha_nacimiento=date(2013, 5, 1), ci="ACTIVO-DEUDA",
        estado="Activo", id_categoria=categoria.id_categoria,
    )
    db.session.add(activo)
    db.session.commit()

    assert retirado.estado_mensualidad == "Retirado"
    respuesta = client.get("/alumnos/?estado=&estado_mensualidad=Con+deuda")
    contenido = respuesta.get_data(as_text=True)
    assert "ActivoDeudorUnico" in contenido
    assert "RetiradoUnico" not in contenido


def test_retirado_no_puede_registrar_pago_ni_cambiar_fecha_base(client):
    alumno, _ = alumno_base(ci="RETIRADO-BLOQUEO", estado="Retirado")
    alumno.fecha_retiro = date.today()
    db.session.commit()
    cantidad_inicial = Pago.query.filter_by(id_alumno=alumno.id_alumno).count()

    respuesta_pago = client.post(
        f"/pagos/alumno/{alumno.id_alumno}/nuevo",
        data={
            "mes_pagado": date.today().strftime("%Y-%m"),
            "fecha_pago": date.today().isoformat(),
            "modo_pago": "Efectivo",
        },
    )
    respuesta_fecha = client.post(
        f"/pagos/alumno/{alumno.id_alumno}/fecha-base",
        data={
            "fecha_base_pago": date.today().isoformat(),
            "motivo": "Intento no permitido",
        },
    )

    assert respuesta_pago.status_code == 302
    assert respuesta_fecha.status_code == 302
    assert Pago.query.filter_by(id_alumno=alumno.id_alumno).count() == cantidad_inicial
    assert alumno.fecha_base_pago is None


def test_partido_campeonato_no_se_edita_desde_ruta_general(client):
    campeonato, categoria, equipos = campeonato_con_equipos(2)
    client.post(f"/campeonatos/{campeonato.id_campeonato}/generar-fixture")
    partido = Partido.query.filter_by(id_campeonato=campeonato.id_campeonato).one()
    respuesta = client.get(f"/partidos/{partido.id_partido}/editar")
    assert respuesta.status_code == 302
    assert f"/campeonatos/{campeonato.id_campeonato}/partidos/" in respuesta.location


def test_partido_campeonato_no_se_elimina_desde_ruta_general(client):
    campeonato, _, _ = campeonato_con_equipos(2)
    client.post(f"/campeonatos/{campeonato.id_campeonato}/generar-fixture")
    partido = Partido.query.filter_by(id_campeonato=campeonato.id_campeonato).one()
    client.post(f"/partidos/{partido.id_partido}/eliminar")
    assert db.session.get(Partido, partido.id_partido) is not None


def test_correccion_resultado_exige_motivo(client):
    campeonato, _, _ = campeonato_con_equipos(2)
    client.post(f"/campeonatos/{campeonato.id_campeonato}/generar-fixture")
    partido = Partido.query.filter_by(id_campeonato=campeonato.id_campeonato).one()
    client.post(
        f"/campeonatos/partidos/{partido.id_partido}/resultado",
        data={"goles_local": "2", "goles_visita": "0"},
    )
    client.post(
        f"/campeonatos/partidos/{partido.id_partido}/resultado",
        data={"goles_local": "0", "goles_visita": "2"},
    )
    assert partido.resultado == "2-0"


def test_numero_qr_activo_no_se_repite(client):
    alumno, categoria = alumno_base(ci="QR-UNO")
    segundo = Alumno(
        nombre="Segundo", apellido="QR", fecha_nacimiento=date(2013, 1, 1),
        ci="QR-DOS", id_categoria=categoria.id_categoria,
    )
    db.session.add(segundo)
    db.session.commit()
    for id_alumno, mes in ((alumno.id_alumno, "2026-06"), (segundo.id_alumno, "2026-07")):
        client.post(
            f"/pagos/alumno/{id_alumno}/nuevo",
            data={
                "mes_pagado": mes, "fecha_pago": "2026-07-20",
                "modo_pago": "QR", "numero_operacion": "OPERACION-UNICA",
            },
        )
    assert Pago.query.filter_by(numero_operacion="OPERACION-UNICA").count() == 1


def test_caja_cerrada_bloquea_nuevo_pago(client):
    alumno, _ = alumno_base()
    db.session.add(CierreCaja(
        fecha=date.today(), efectivo_sistema=0, qr_sistema=0,
        efectivo_contado=0, qr_confirmado=0, diferencia=0,
    ))
    db.session.commit()
    client.post(
        f"/pagos/alumno/{alumno.id_alumno}/nuevo",
        data={
            "mes_pagado": date.today().strftime("%Y-%m"),
            "fecha_pago": date.today().isoformat(), "modo_pago": "Efectivo",
        },
    )
    assert Pago.query.count() == 0


def test_calendario_rechaza_mes_invalido(client):
    respuesta = client.get("/calendario/?anio=2026&mes=13")
    assert respuesta.status_code == 302


def test_calendario_muestra_mes_en_espanol_y_todas_las_actividades(client):
    from app.routes.calendario import MESES_ES
    hoy = date.today()
    for indice in range(1, 7):
        db.session.add(EventoAgenda(
            fecha=hoy, descripcion=f"Actividad completa {indice}",
            tipo="Evento",
        ))
    db.session.commit()

    respuesta = client.get(f"/calendario/?anio={hoy.year}&mes={hoy.month}")

    assert respuesta.status_code == 200
    assert f"{MESES_ES[hoy.month - 1]} {hoy.year}".encode() in respuesta.data
    assert b"Agenda desde hoy" in respuesta.data
    assert b"PDF semanal" in respuesta.data
    assert b'/calendario/exportar-semanal.pdf' in respuesta.data
    for indice in range(1, 7):
        assert f"Actividad completa {indice}".encode() in respuesta.data


def test_calendario_filtra_campeonatos_por_categoria(client):
    categoria_visible = Categoria.query.filter_by(nombre="13").first()
    categoria_oculta = Categoria.query.filter_by(nombre="11").first()
    db.session.add_all([
        Campeonato(nombre="Copa visible", fecha_inicio=date(2026, 8, 10), categorias=[categoria_visible]),
        Campeonato(nombre="Copa de otra categorÃ­a", fecha_inicio=date(2026, 8, 10), categorias=[categoria_oculta]),
    ])
    db.session.commit()

    respuesta = client.get(
        "/calendario/",
        query_string={"anio": 2026, "mes": 8, "id_categoria": categoria_visible.id_categoria},
    )

    assert b"Copa visible" in respuesta.data
    assert "Copa de otra categorÃ­a".encode() not in respuesta.data


def test_asistencia_guardada_aparece_en_perfil_del_alumno(client):
    alumno, categoria = alumno_base(ci="ASISTENCIA-PERFIL")
    equipo = Equipo(nombre="Equipo del perfil", id_categoria=categoria.id_categoria)
    alumno.equipos.append(equipo)
    entrenamiento = Entrenamiento(
        dia_semana=2, hora_inicio=time(18, 0), hora_fin=time(19, 30),
        lugar="Cancha de asistencia", categoria=categoria, equipo=equipo,
    )
    sesion = SesionEntrenamiento(
        entrenamiento=entrenamiento, fecha=date(2026, 8, 19), estado="Realizada",
    )
    db.session.add(AsistenciaEntrenamiento(
        sesion=sesion, alumno=alumno, estado="Presente",
    ))
    db.session.commit()

    respuesta = client.get(f"/alumnos/{alumno.id_alumno}")

    assert respuesta.status_code == 200
    assert b"Asistencia y programaci" in respuesta.data
    assert b"19/08/2026" in respuesta.data
    assert b"Cancha de asistencia" in respuesta.data
    assert b"Equipo del perfil" in respuesta.data
    assert b"100%" in respuesta.data


def test_turno_alumno_se_persiste_y_aparece_en_perfil(client):
    respuesta = client.post(
        "/alumnos/nuevo",
        data={
            "nombre": "Lucía", "apellido": "Turno", "fecha_nacimiento": "2014-05-10",
            "estado": "Activo", "beca": "Ninguna", "turno": "Tarde",
        },
        follow_redirects=True,
    )

    alumno = Alumno.query.filter_by(nombre="Lucía", apellido="Turno").one()
    assert respuesta.status_code == 200
    assert alumno.turno == "Tarde"
    assert b"Turno de entrenamiento" in respuesta.data
    assert b"Tarde" in respuesta.data


def test_asistencia_respeta_turno_persistido_del_alumno(client):
    alumno_manana, categoria = alumno_base(ci="TURNO-MANANA")
    alumno_manana.nombre = "Alumno Mañana"
    alumno_manana.turno = "Mañana"
    alumno_noche = Alumno(
        nombre="Alumno Noche", apellido="Control", fecha_nacimiento=date(2013, 5, 1),
        ci="TURNO-NOCHE", estado="Activo", id_categoria=categoria.id_categoria, turno="Noche",
    )
    entrenamiento = Entrenamiento(
        dia_semana=1, hora_inicio=time(9, 0), hora_fin=time(10, 30), categoria=categoria,
    )
    db.session.add_all([alumno_noche, entrenamiento])
    db.session.commit()

    respuesta = client.get(f"/calendario/entrenamientos/{entrenamiento.id_entrenamiento}/asistencia")

    assert respuesta.status_code == 200
    assert "Alumno Mañana".encode() in respuesta.data
    assert b"Alumno Noche" not in respuesta.data


def test_consulta_individual_de_asistencia_filtra_ultimas_tres(client):
    alumno, categoria = alumno_base(ci="CONSULTA-ASISTENCIA")
    alumno.turno = "Tarde"
    entrenamiento = Entrenamiento(
        dia_semana=0, hora_inicio=time(15, 0), hora_fin=time(16, 30), categoria=categoria,
    )
    db.session.add(entrenamiento)
    for indice, estado in enumerate(["Presente", "Ausente", "Tarde", "Justificado", "Presente"]):
        sesion = SesionEntrenamiento(
            entrenamiento=entrenamiento, fecha=date.today() - timedelta(days=indice), estado="Realizada",
        )
        db.session.add(AsistenciaEntrenamiento(sesion=sesion, alumno=alumno, estado=estado))
    db.session.commit()

    respuesta = client.get(
        "/alumnos/asistencias",
        query_string={"id_alumno": alumno.id_alumno, "periodo": "ultimas3"},
    )

    assert respuesta.status_code == 200
    assert b"Asistencia de jugadores" in respuesta.data
    assert alumno.nombre_completo.encode() in respuesta.data
    assert b"3 registros" in respuesta.data
    assert b"Justificado" in respuesta.data


def test_consulta_asistencia_filtra_estado_y_turno(client):
    alumno, categoria = alumno_base(ci="FILTRO-ASISTENCIA")
    alumno.turno = "Noche"
    entrenamiento = Entrenamiento(
        dia_semana=0, hora_inicio=time(19, 0), hora_fin=time(20, 0), categoria=categoria,
    )
    sesion = SesionEntrenamiento(entrenamiento=entrenamiento, fecha=date.today(), estado="Realizada")
    db.session.add(AsistenciaEntrenamiento(sesion=sesion, alumno=alumno, estado="Ausente"))
    db.session.commit()

    respuesta = client.get(
        "/alumnos/asistencias",
        query_string={"periodo": "todas", "turno": "Noche", "estado_asistencia": "Ausente"},
    )

    assert respuesta.status_code == 200
    assert alumno.nombre_completo.encode() in respuesta.data
    assert b"Ausente" in respuesta.data


def test_agenda_rechaza_equipo_inexistente(client):
    respuesta = client.post(
        "/agenda/nuevo",
        data={
            "fecha": date.today().isoformat(), "descripcion": "Evento inválido",
            "id_equipo": "999999",
        },
    )
    assert respuesta.status_code == 200
    assert EventoAgenda.query.filter_by(descripcion="Evento inválido").count() == 0
