import io
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.utils import get_column_letter

from app.utils.marca_agua import imagen_con_opacidad


def valor_seguro_excel(valor):
    """Impide que datos ingresados por usuarios se ejecuten como fórmulas."""
    if isinstance(valor, str) and valor.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + valor
    return valor


def _es_columna_numeracion(encabezado) -> bool:
    texto = str(encabezado or "").strip().lower()
    return texto.startswith("n") and len(texto) <= 4


def _titulo_hoja_seguro(titulo: str) -> str:
    """Limpia los caracteres que Excel no admite en nombres de hojas."""
    limpio = re.sub(r"[\\/*?:\[\]]+", " ", str(titulo or "Datos"))
    limpio = " ".join(limpio.split()).strip("' ") or "Datos"
    return limpio[:31]


def exportar_tabla_excel(titulo: str, encabezados: list, filas: list, logo_path=None, orientacion="vertical") -> io.BytesIO:
    """Exportador genérico a Excel para listados con encabezados y filas ya
    formateados como texto (usado por Alumnos, Partidos, etc.)."""
    wb = Workbook()
    ws = wb.active
    ws.title = _titulo_hoja_seguro(titulo)
    wb.properties.title = str(titulo or "Datos")[:255]
    wb.properties.creator = "Sistema de Gestión de Academias"
    ws.sheet_view.showGridLines = False

    fila_encabezado = 1
    borde_inferior = Border(bottom=Side(style="thin", color="D5DED7"))

    for indice, encabezado in enumerate(encabezados, start=1):
        ws.cell(row=fila_encabezado, column=indice, value=encabezado)
    for col in range(1, len(encabezados) + 1):
        celda = ws.cell(row=fila_encabezado, column=col)
        celda.font = Font(bold=True, color="000000")
        celda.fill = PatternFill("solid", fgColor="2E7D32")
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[fila_encabezado].height = 30

    for numero_fila, fila in enumerate(filas, start=fila_encabezado + 1):
        for numero_columna, valor in enumerate(fila, start=1):
            celda = ws.cell(
                row=numero_fila, column=numero_columna,
                value=valor_seguro_excel(valor),
            )
            celda.alignment = Alignment(vertical="top", wrap_text=True)
            celda.border = borde_inferior
            if isinstance(valor, str) and valor.lower().startswith(("http://", "https://")):
                celda.hyperlink = valor
                celda.style = "Hyperlink"

    if logo_path:
        marca = ExcelImage(imagen_con_opacidad(logo_path, 0.08))
        proporcion = marca.width / marca.height
        marca.height = 300
        marca.width = 300 * proporcion
        columna_central = max(1, (len(encabezados) + 1) // 2)
        fila_central = max(2, fila_encabezado + max(1, len(filas) // 2) - 8)
        ws.add_image(marca, f"{get_column_letter(columna_central)}{fila_central}")

    ws.freeze_panes = f"A{fila_encabezado + 1}"
    if encabezados:
        ws.auto_filter.ref = f"A{fila_encabezado}:{get_column_letter(len(encabezados))}{max(fila_encabezado, len(filas) + 1)}"
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.print_title_rows = "1:1"
    if orientacion == "horizontal":
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    else:
        ws.page_setup.orientation = "portrait"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    for indice_columna, col_cells in enumerate(ws.columns, start=1):
        largo = max(len(str(c.value)) if c.value else 0 for c in col_cells)
        letra = col_cells[0].column_letter
        if indice_columna == 1 and encabezados and _es_columna_numeracion(encabezados[0]):
            # Sólo el espacio necesario para "N°" y la cantidad de dígitos.
            ws.column_dimensions[letra].width = max(3.5, min(7, largo + 0.8))
            for celda in col_cells:
                celda.alignment = Alignment(horizontal="center")
        else:
            ws.column_dimensions[letra].width = max(8, min(36, largo + 3))

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def exportar_pagos_excel(pagos) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Pagos"

    encabezados = [
        "ID Pago", "Alumno", "CI", "Categoría", "Mes pagado",
        "Fecha de pago", "Hora de registro", "Fecha de vencimiento",
        "Estado", "Modo de pago", "Beca", "Monto",
    ]
    ws.append(encabezados)
    for col in range(1, len(encabezados) + 1):
        celda = ws.cell(row=1, column=col)
        celda.font = Font(bold=True, color="000000")
        celda.fill = PatternFill("solid", fgColor="2E7D32")
        celda.alignment = Alignment(horizontal="center")

    for p in pagos:
        ws.append([valor_seguro_excel(valor) for valor in [
            p.id_pago,
            p.alumno.nombre_completo,
            p.alumno.ci or "",
            p.alumno.categoria.nombre if p.alumno.categoria else "",
            p.mes_pagado,
            p.fecha_pago.strftime("%d/%m/%Y"),
            p.hora_registro.strftime("%H:%M") if p.hora_registro else "",
            p.fecha_vencimiento.strftime("%d/%m/%Y"),
            p.estado,
            "Beca" if p.es_beca else p.modo_pago,
            p.alumno.beca,
            p.monto,
        ]])

    for col_cells in ws.columns:
        largo = max(len(str(c.value)) if c.value else 0 for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = largo + 3

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
